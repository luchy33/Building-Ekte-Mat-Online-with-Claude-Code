"""Build and validate EkteMat_backlog_v1.xlsx.

Purpose (INFO212, Ekte Mat Online):
  1. Hold the backlog content (stakeholders, stories, NFRs, issues) as data.
  2. Verify that every supporting quote occurs verbatim in the cited source file
     (whitespace-normalised) or, for the CSV, in the body of the cited email_id.
  3. Check ID uniqueness/sequence, epic sizes and per-file story counts.
  4. Write the workbook with Excel tables, filters, frozen headers, wrapping and
     data validation. Audit counts are live COUNTIF formulas.

Usage: python -I EkteMat_backlog_v1_build.py <repo_root_copy> <output_xlsx>
Source files are only read, never modified.
"""
import csv, re, sys, collections
from pathlib import Path
import docx
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

REPO = Path(sys.argv[1]); OUT = Path(sys.argv[2])

SRC = {  # key -> path relative to EkteMat_repo
 "OWN": "01_Requirements/sources/01_interview_owner_2026-08-14.txt",
 "EVC": "01_Requirements/sources/01_interview_event-coordinator_2026-08-15.txt",
 "CHF": "01_Requirements/sources/01_interview_head-chef_2026-08-15.txt",
 "CORP": "01_Requirements/sources/01_interview_corporate-customer_2026-08-19.txt",
 "GDPR": "01_Requirements/sources/01_note_gdpr-consultant_2026-08-20.txt",
 "VIS": "01_Requirements/sources/01_vision_draft-v0.3_2026-10-20.docx",
 "MIN": "00_Project_management/00_minutes_kickoff-meeting_2026-08-12.txt",
 "BUD": "00_Project_management/00_plan_budget-and-constraints_2026-10-20.txt",
 "ASIS": "02_Analysis_and_models/02_process_as-is-order-handling_2026-08-15.md",
 "ITL": "03_Architecture/03_landscape_current-it_2026-10-20.txt",
 "CSV": "04_Data/04_dataset_customer-emails_2026-10-20.csv",
 "MENU": "04_Data/04_reference_lunch-menu-autumn_2026-10-20.txt",
}
norm = lambda s: re.sub(r"\s+", " ", s).strip()
TEXT, EMAIL = {}, {}
for k, rel in SRC.items():
    p = REPO / rel
    if k == "VIS":
        TEXT[k] = norm("\n".join(par.text for par in docx.Document(p).paragraphs))
    elif k == "CSV":
        rows = list(csv.DictReader(open(p, encoding="utf-8")))
        EMAIL = {r["email_id"]: norm(r["body"]) for r in rows}
        TEXT[k] = ""
    else:
        TEXT[k] = norm(p.read_text(encoding="utf-8"))

def Q(key, ref, text, tr=None):
    return (key, ref, text, tr)

def render(quotes):
    out = []
    for k, ref, t, tr in quotes:
        s = f'{Path(SRC[k]).name} [{ref}]: "{t}"'
        if tr: s += f" [translation: {tr}]"
        out.append(s)
    return "\n".join(out)

def sources(quotes):
    seen = []
    for k, *_ in quotes:
        if SRC[k] not in seen: seen.append(SRC[k])
    return "\n".join(seen)

# ---------------------------------------------------------------- stakeholders
STAKE = [
 ("Ingrid Solheim", "Founder, owner (daglig leder) and product owner",
  "Stop losing inquiries; self-service menus and booking; kitchen overview; deposit; keep Fiken; show food origin; live by 2027-03-01",
  "High", "Lost inquiries (about 60 000 NOK); no-shows at weddings; budget about 450 000 NOK; deadline",
  [Q("OWN","header","Ingrid Solheim (founder and owner, Ekte Mat AS)"),
   Q("MIN","bullet 5","Ingrid acts as product owner; available Tuesdays 09-10")],
  "DIRECT", "Influence High: owner and appointed product owner (documented)."),
 ("Sunniva Berg", "Event coordinator",
  "One order record instead of Excel + printouts; capacity calendar; automatic 7-day reminders; allergy confirmation",
  "Medium", "Forgotten updates; slow replies in peak season; overbooking; corporate customers lost if deposit is forced",
  [Q("EVC","header","Sunniva Berg (event coordinator, Ekte Mat AS)"),
   Q("EVC","Q3","If we force a deposit on them, we will lose them.")],
  "ASSUMPTION", "Influence Medium is an assessment: key daily user, no documented decision authority."),
 ("Mateo Rinaldi", "Head chef",
  "Kitchen wall-tablet overview; PIN login; edit seasonal menus; supplier-linked ingredients; prep list; offline access",
  "Medium", "Allergy visibility; speed; internet outages",
  [Q("CHF","header","Mateo Rinaldi (head chef)"),
   Q("CHF","Q1","One screen in the kitchen. Tablet on the wall.")],
  "ASSUMPTION", "Influence Medium is an assessment: owns kitchen operations, no documented project authority."),
 ("Henrik Aas (Fjordtek AS)", "Office manager at Ekte Mat's biggest corporate customer",
  "Company login; standing orders; copy orders; change until 10:00 day before; approval above 15 000 NOK; monthly EHF invoice",
  "Medium", "Late confirmations; no prepayment allowed; EU/EEA data storage; deletion of employee names",
  [Q("CORP","header","Henrik Aas (office manager, Fjordtek AS - fictional company, 180 employees)"),
   Q("OWN","Q6","maybe Fjordtek - they are our biggest corporate customer.")],
  "ASSUMPTION", "Influence Medium is an assessment based on 'biggest corporate customer'."),
 ("Fjordtek IT department", "IT function of a corporate customer (not interviewed)",
  "Data location and deletion", "Low", "Data stored in EU/EEA; employee names deleted at end of relationship",
  [Q("CORP","Q4","Our IT department will ask where our data is stored.")],
  "ASSUMPTION", "Group named only indirectly; influence assessed."),
 ("Fjordtek finance department", "Finance function of a corporate customer (not interviewed)",
  "Invoice-based payment", "Low", "Prepayment not allowed",
  [Q("CORP","Q3","We cannot prepay, our finance department does not allow it.")],
  "ASSUMPTION", "Group named only indirectly; influence assessed."),
 ("Corporate customers (group)", "Customer segment: weekly lunches and meetings",
  "Recurring orders, several orderers, EHF invoicing with references, English menus", "Medium",
  "PDF invoices rejected; repetitive weekly e-mails; reorder without details",
  [Q("MIN","bullet 3","Customer groups: corporate (weekly lunches, meetings), private (parties), weddings, health-conscious"),
   Q("CSV","E004","our finance department only accepts EHF invoices with our PO number."),
   Q("CSV","E007","It is a pain to e-mail every week.")],
  "ASSUMPTION", "Segment documented; influence assessed."),
 ("Private customers (parties)", "Customer segment: parties, birthdays, confirmations",
  "Vipps payment, e-mail receipt, pick-up option, phone ordering", "Low", "Payment method; ability to order without internet",
  [Q("MIN","bullet 3","private (parties)"),
   Q("CSV","E009","Kan jeg betale med Vipps for bursdag 30 pers?","Can I pay with Vipps for a birthday party of 30?"),
   Q("CSV","E024","Kan jeg fortsatt bestille på telefon?","Can I still order by phone?")],
  "ASSUMPTION", "Segment documented; influence assessed."),
 ("Wedding customers", "Customer segment: weddings (season from May 2027)",
  "Capacity for served dinners; deposit amount", "Low", "Availability on the date; deposit terms",
  [Q("OWN","Q2","We have had too many no-shows at weddings."),
   Q("CSV","E005","Har dere kapasitet? Hvor mye må vi betale i depositum?","Do you have capacity? How much do we have to pay as a deposit?")],
  "ASSUMPTION", "Segment documented; influence assessed."),
 ("Health-conscious customers", "Customer segment named at kick-off",
  "Not documented beyond the segment name", "Low", "Not documented (no interview or e-mail attributable)",
  [Q("MIN","bullet 3","health-conscious")],
  "ASSUMPTION", "Only the segment name is evidenced; concerns unknown."),
 ("Driver (name not documented)", "Delivers orders (not interviewed)",
  "Not documented", "Low", "Currently gets the address from a printout and uses own phone",
  [Q("ASIS","step 10","Kitchen prepares, driver delivers. Driver gets the address from the printout."),
   Q("ITL","Driver","Driver: own phone, Google Maps")],
  "ASSUMPTION", "Role evidenced; needs not elicited (see issue register)."),
 ("Data-protection consultant (name not documented)", "External data-protection adviser",
  "GDPR obligations", "Medium", "Health data, retention, data-subject rights, data processing agreements",
  [Q("GDPR","header","Note from data-protection consultant (phone call, 20 Aug 2026)")],
  "ASSUMPTION", "Advisory role; influence assessed."),
 ("Student project team", "Three student developers, 2-week sprints",
  "Deliver first version by go-live", "Medium", "Budget, deadline, scope",
  [Q("BUD","Team","Team: 3 student developers, 2-week sprints, product owner Ingrid (Tue 09-10)")],
  "ASSUMPTION", "Influence assessed: delivers the system but does not own scope."),
 ("Suppliers (farms, boats)", "Ingredient suppliers whose origin is shown",
  "Being shown as origin; becoming a supplier", "Low", "Not elicited",
  [Q("OWN","Q4","where the lamb or the fish comes from - the farm, the boat."),
   Q("CHF","Q2","ingredients should be linked to suppliers"),
   Q("CSV","E025","we are an organic vegetable farm in Hardanger and would like to become a supplier.")],
  "ASSUMPTION", "Group evidenced; influence assessed."),
]

# --------------------------------------------------------------------- epics
EPICS = {
 "EP-01": "Online menus and booking requests",
 "EP-02": "Order management and event coordination",
 "EP-03": "Kitchen operations",
 "EP-04": "Payments and invoicing",
 "EP-05": "Corporate customer accounts",
 "EP-06": "Privacy and data-subject rights",
}
PO = "PO REVIEW: "
# (epic, story, [AC], moscow, rationale, quotes, status, assumptions)
ST = [
 ("EP-01","As a prospective customer, I want to see current menus and prices online at any time so that I can plan an event without e-mailing for a quote.",
  ["Given a published menu, when a visitor opens the menu page, then each dish is shown with its price in NOK.",
   "Given it is outside office hours, when a visitor opens the website, then menus and prices are available without contacting staff."],
  "Must","Named by the owner as a first function; core of the documented project goal.",
  [Q("OWN","Q2","Customers must be able to see our menus and prices and send a booking request themselves, at any hour."),
   Q("CSV","E001","Can you send the menu and price?"),
   Q("CSV","E015","Can you send a price list PDF?"),
   Q("VIS","para 1","lets customers see menus, prices and food origin")],
  "DIRECT","Pricing unit (per person, as in the autumn menu draft) assumed for all packages; VAT display and minimum order not documented."),
 ("EP-01","As a prospective customer, I want to send a booking request online at any hour so that my request is registered and not lost.",
  ["Given a visitor has completed the booking form, when they submit it, then the request is stored and listed for the coordinator.",
   "Given the submission cannot be stored, when the visitor submits, then an error message is shown to the visitor."],
  "Must","Lost inquiries are the main documented problem; replaces e-mail intake (kick-off goal).",
  [Q("OWN","Q1","We lose inquiries. Last December we found three julebord requests in the spam folder after the date had passed."),
   Q("MIN","bullet 1","Goal: replace e-mail/phone/Excel ordering with an online ordering and event booking system"),
   Q("CSV","E023","I tried the contact form on your website three times and got an error."),
   Q("CSV","E006","We have booked another caterer now.")],
  "DIRECT","Required booking fields are not defined. Whether a request is auto-confirmed or reviewed by the coordinator is not decided."),
 ("EP-01","As a prospective customer, I want to see which dates are available so that I only request dates Ekte Mat can serve.",
  ["Given a date has reached the capacity limit, when a customer views the availability calendar, then the date is shown as unavailable.",
   "Given a date has remaining capacity, when a customer views the calendar, then the date can be selected for a request."],
  "Should","Clear customer need, but booking requests can be handled without it at go-live.",
  [Q("CSV","E001","Is it possible to see online which dates are free?"),
   Q("CSV","E005","Har dere kapasitet?","Do you have capacity?")],
  "DIRECT","Depends on US-10 and the capacity rule (ISS-06). Whether exact remaining capacity is shown publicly is not decided."),
 ("EP-01","As a customer, I want an automatic confirmation and a visible status for my order so that I know whether it is accepted.",
  ["Given a booking request is stored, when submission completes, then the customer receives a message that the request was received.",
   "Given the coordinator confirms the order, when the status changes, then the customer is notified that the order is confirmed."],
  "Should","Strongly evidenced pain point, but not named as a go-live condition.",
  [Q("CORP","Q1","Then I wait for a confirmation. Sometimes it arrives Monday morning, which is too late for me."),
   Q("CSV","E011","Is our order for tomorrow confirmed? We never got a confirmation e-mail.")],
  "DIRECT",PO+"priority could be Must because it is closely tied to the lost-inquiry problem. Notification channel (e-mail) is an ASSUMPTION."),
 ("EP-01","As a customer, I want menus to show where key ingredients come from (farm or boat) so that I can choose food with a known origin.",
  ["Given a dish has an ingredient linked to a supplier, when a customer views the menu, then the origin of that ingredient is shown."],
  "Should","Part of the vision and brand, but stated by the owner as a wish ('would love').",
  [Q("OWN","Q4","I would love every menu to show where the lamb or the fish comes from - the farm, the boat."),
   Q("CSV","E008","Which farm is the lamb from?"),
   Q("VIS","para 1","see menus, prices and food origin")],
  "DIRECT","Depends on US-19. Which ingredients must show origin ('every menu' vs key ingredients) is not defined. Only one AC is supported by evidence."),
 ("EP-01","As a customer, I want the allergens in each dish to be clearly marked so that guests with allergies can choose safely.",
  ["Given a dish contains an allergen on the labelling list, when the menu is displayed, then that allergen is shown for the dish.",
   "Given the English version of the menu, when it is displayed, then the same allergen markings are shown."],
  "Should","Safety-relevant and requested; no source states it as a go-live condition.",
  [Q("MENU","last line","Allergens to be labelled: fish, milk, gluten, nuts, celery, mustard."),
   Q("CSV","E012","Is your menu available in English with allergens marked?"),
   Q("CSV","E019","Please label gluten-free items clearly on the buffet.")],
  "DIRECT",PO+"consider Must. Legal allergen-labelling obligations are not in the evidence. E019 concerns physical buffet labels, so applying it to the online menu is INFERRED. Depends on NFR-10."),
 ("EP-01","As a customer, I want to record allergies and diets per guest group in my order so that the kitchen can prepare safe meals.",
  ["Given an order, when the customer adds a guest group, then they can enter the number of guests and their allergies or diets for that group.",
   "Given allergy information is entered, when the order is saved, then no guest name is required for the allergy data."],
  "Must","Documented as mandatory for every order; safety incident in 2024; GDPR advice on health data.",
  [Q("EVC","Q5","Every order has to show allergies and diets per guest group"),
   Q("GDPR","bullet 1","Store per guest group, not per named guest, unless strictly needed."),
   Q("CSV","E002","To av gjestene har nøtteallergi og en er cøliaki.","Two of the guests have a nut allergy and one has coeliac disease."),
   Q("CSV","E022","Adding one guest with shellfish allergy to the order for tonight.")],
  "DIRECT","Standard list of allergens/diets: the menu draft lists six allergens; whether the order form uses the same list is not decided."),
 ("EP-01","As a corporate customer, I want to see the share of local ingredients per order so that I can use it in sustainability (ESG) reporting.",
  ["Given an order whose dishes have ingredients with recorded origin, when the customer views the order, then the share of local ingredients is shown."],
  "Could","Single customer request; not linked to go-live needs.",
  [Q("CSV","E021","can you tell us the CO2 footprint or at least the share of local ingredients per order?")],
  "DIRECT","Definition of 'local' and the calculation basis (weight, cost, count) are missing. Depends on US-19. CO2 footprint not included (no data source)."),
 ("EP-02","As the event coordinator, I want to register and manage all orders, including orders received by phone or e-mail, in one system so that I no longer keep Excel and printouts in sync.",
  ["Given an order received by phone, when the coordinator registers it, then it appears in the same order list as online orders.",
   "Given the coordinator changes an order, when the change is saved, then the kitchen overview shows the updated order without re-printing."],
  "Must","Replacing e-mail/phone/Excel ordering is the documented project goal.",
  [Q("EVC","Q1","I have to remember to update both the sheet and the printout. Sometimes I forget."),
   Q("ASIS","step 1","Customer sends an e-mail or calls (phone calls are noted on paper)."),
   Q("CSV","E024","Kan jeg fortsatt bestille på telefon?","Can I still order by phone?"),
   Q("VIS","para 1","it gives the kitchen one reliable overview of every order and allergy.")],
  "DIRECT","Assumes phone remains an order channel (ISS-12). Migration of existing Excel orders is not addressed."),
 ("EP-02","As the event coordinator, I want a calendar showing booked guests per day against kitchen capacity so that I do not overbook.",
  ["Given confirmed orders on a date, when the coordinator opens the calendar, then the total number of guests for that date is shown.",
   "Given the booked guests for a date exceed the capacity limit, when the calendar is shown, then the date is flagged as over capacity."],
  "Should","Clear operational need; booking requests are still reviewed manually.",
  [Q("EVC","Q4","A calendar that shows how many guests we already have per day, so I do not overbook. Our kitchen can do about 400 covers a day, 250 if there is a wedding with a served dinner."),
   Q("ASIS","step 3","If the day is \"full\" (gut feeling), she proposes another date.")],
  "DIRECT",PO+"becomes Must if online requests are confirmed automatically. Capacity values are approximate ('about'): see ISS-06."),
 ("EP-02","As a customer, I want to change the number of guests and allergy information for my order online so that changes reach Ekte Mat without phone calls.",
  ["Given an existing order before the change deadline, when the customer changes the number of guests, then the order shows the new number for the coordinator and the kitchen.",
   "Given a customer adds a guest group with an allergy, when the change is saved, then the kitchen overview shows the allergy for that order."],
  "Should","Frequent pain point; not documented as a go-live condition.",
  [Q("CORP","Q2","change the number of people until 10:00 the day before"),
   Q("CSV","E003","we need to change from 40 to 55 guests for our event next Thursday. I called yesterday but nobody picked up."),
   Q("ASIS","step 8","Changes (guest count, allergies) arrive by phone/e-mail"),
   Q("VIS","para 1","book and change orders any time")],
  "DIRECT","The change deadline is not agreed: 10:00 the day before is one customer's wish, while the vision says 'any time' (ISS-03)."),
 ("EP-02","As the event coordinator, I want customers to get an automatic reminder seven days before the event to confirm final guest numbers so that I do not have to phone everyone.",
  ["Given an event is seven days away, when the daily reminder job runs, then the customer receives a reminder to confirm final guest numbers.",
   "Given a reminder has been sent, when the coordinator opens the order, then the order shows that the reminder was sent."],
  "Should","Reduces manual work; current manual step exists as a fallback.",
  [Q("EVC","Q4","automatic reminders to customers about final guest numbers seven days before the event."),
   Q("ASIS","step 9","7 days before the event the coordinator phones the customer to confirm final numbers (if she has time).")],
  "DIRECT","Reminder channel and what happens if the customer does not respond are not documented."),
 ("EP-02","As a customer, I want to request cancellation of an order so that Ekte Mat is informed in time.",
  ["Given an existing order, when the customer requests cancellation, then the order status shows 'cancellation requested' and the coordinator is notified."],
  "Could","Single request; cancellation terms are undefined.",
  [Q("CSV","E013","Unfortunately we have to cancel the dinner on 3 October")],
  "UNRESOLVED","Cancellation terms and fees are missing (ISS-04). The AC is an ASSUMPTION about the mechanism."),
 ("EP-02","As a customer, I want to choose pick-up or delivery and state the delivery time so that the food arrives when and where I need it.",
  ["Given a customer selects delivery, when they submit the order, then a delivery address and a requested delivery time are required.",
   "Given a customer selects pick-up, when they submit the order, then no delivery address is required."],
  "Should","Needed for fulfilment, but can be handled through notes at go-live.",
  [Q("CSV","E014","Kan vi hente maten selv i stedet for levering?","Can we collect the food ourselves instead of having it delivered?"),
   Q("CSV","E018","Can you deliver between 10:45 and 11:00 exactly?"),
   Q("ASIS","step 10","Driver gets the address from the printout.")],
  "DIRECT","Delivery area, fees and time-window precision are not documented. Driver needs are not elicited (ISS-13)."),
 ("EP-03","As the head chef, I want a daily overview on a kitchen wall tablet showing orders, portions per dish and allergies in red so that the kitchen prepares the right food safely.",
  ["Given confirmed orders for a day, when the chef opens the day view, then each order and the total portions of each dish are shown.",
   "Given an order contains allergy or diet information, when it is displayed, then that information is shown in red."],
  "Must","The owner and the vision require a kitchen order overview; it replaces printouts.",
  [Q("CHF","Q1","For each day: which orders, how many portions of each dish, allergies in red."),
   Q("OWN","Q2","I want a proper order overview for the kitchen"),
   Q("ASIS","step 7","Coordinator prints the order and pins it on the kitchen board."),
   Q("CSV","E017","The note said 'V x4' but nothing came.")],
  "DIRECT","Depends on NFR-03 (offline) and NFR-13 (device). E017 supports clear diet visibility (INFERRED)."),
 ("EP-03","As the event coordinator, I want the kitchen to confirm that it has seen each order's allergy information so that allergies are not missed.",
  ["Given an order with allergy information, when a kitchen user confirms it, then the confirmation is recorded with time and user.",
   "Given an order's allergy information has not been confirmed, when the coordinator views the order, then it is marked as unconfirmed."],
  "Must","Stated as mandatory after a near-miss nut incident in 2024; customers ask for this assurance.",
  [Q("EVC","Q5","the kitchen must confirm they have seen it. We had a nut incident in 2024 - nobody was hurt, but it was close."),
   Q("CSV","E002","Kan dere bekrefte at kjøkkenet har fått beskjed?","Can you confirm that the kitchen has been informed?")],
  "DIRECT",PO+"Must is based on safety evidence, not on a stated go-live condition. Whether a later allergy change resets the confirmation is an ASSUMPTION to confirm."),
 ("EP-03","As a kitchen staff member, I want to log in to the kitchen tablet with a PIN so that I can use it quickly while working.",
  ["Given the kitchen tablet, when a user enters a valid PIN, then the kitchen view opens.",
   "Given an invalid PIN is entered, when the user submits it, then access is denied."],
  "Should","Usability need for the kitchen view; not a go-live condition on its own.",
  [Q("CHF","Q1","I do not want to log in with long passwords with flour on my hands - a PIN is fine.")],
  "DIRECT","PIN length, lock-out rules and the security of health data access are not documented (ISS-14)."),
 ("EP-03","As the head chef, I want to edit the menu myself so that seasonal changes are published without help.",
  ["Given the chef is logged in with menu rights, when he changes a dish and saves, then the published menu shows the change."],
  "Should","Menus change four times a year; the first change can be handled by the team.",
  [Q("CHF","Q2","Menus change with the season, four times a year. I want to edit the menu myself"),
   Q("ITL","Constraints","nobody technical in-house")],
  "DIRECT","Who may change prices is not documented. Related to NFR-12."),
 ("EP-03","As the head chef, I want ingredients to be linked to suppliers so that dish origin can be shown to customers.",
  ["Given a supplier is registered, when the chef links an ingredient to it, then dishes using that ingredient show the supplier as origin."],
  "Should","Enables US-05 (Should); not a go-live condition.",
  [Q("CHF","Q2","ingredients should be linked to suppliers, so we can show origin like Ingrid wants.")],
  "DIRECT","Supplier data fields are not defined."),
 ("EP-03","As the head chef, I want a prep list with total kilos of each ingredient for the next three days so that I can plan purchasing and preparation.",
  ["Given confirmed orders for the next three days, when the chef opens the prep list, then the total kilos of each ingredient are shown."],
  "Could","Useful, but depends on recipe quantity data that does not exist in the evidence.",
  [Q("CHF","Q2","I need a prep list: total kilos of each ingredient for the next three days.")],
  "DIRECT","Requires ingredient quantities per portion (missing evidence, ISS-15)."),
 ("EP-04","As a private customer, I want to pay with Vipps and get a receipt by e-mail so that payment is easy and documented.",
  ["Given a private customer's order, when they choose to pay, then Vipps is offered as a payment method.",
   "Given a Vipps payment is completed, when the order is updated, then it is marked as paid and the customer receives a receipt by e-mail."],
  "Must","Kick-off states the system must integrate with Vipps; the owner insists on it.",
  [Q("OWN","Q3","Vipps for private customers, definitely."),
   Q("MIN","bullet 2","Must integrate with Fiken (accounting) and Vipps (payments)."),
   Q("CSV","E009","Kan jeg betale med Vipps for bursdag 30 pers? Og kan jeg få kvittering på e-post?","Can I pay with Vipps for a birthday party of 30? And can I get a receipt by e-mail?")],
  "DIRECT","Vipps agreement and API details are external dependencies (ISS-09). The e-mail receipt is from one customer request."),
 ("EP-04","As the owner, I want a deposit to be required for defined booking types so that no-shows are reduced.",
  ["Given a booking of a type that requires a deposit, when the booking is made, then a deposit request is issued and the booking stays unconfirmed until it is paid."],
  "Should","Scope is in conflict (all bookings vs weddings only); priority depends on the PO decision.",
  [Q("OWN","Q2","I want the system to take a 30 % deposit when someone books - no deposit, no booking."),
   Q("EVC","Q3","Honestly, corporate clients hate prepayment."),
   Q("EVC","Q3","For weddings a deposit makes sense."),
   Q("CORP","Q3","We cannot prepay, our finance department does not allow it."),
   Q("ASIS","step 6","For weddings, the coordinator sends a Vipps payment request for a deposit (sometimes forgotten).")],
  "UNRESOLVED",PO+"see ISS-01. Likely Must for weddings if scope is confirmed: wedding bookings for May 2027 will be made before go-live. Refund rules are not documented."),
 ("EP-04","As a corporate customer, I want one EHF invoice per month with my reference (cost centre or PO number) on each line so that our finance department can process it.",
  ["Given a corporate customer on monthly invoicing, when the month is closed, then one EHF invoice containing all that month's orders is produced.",
   "Given an order has a customer reference, when the invoice is produced, then the reference appears on that invoice line."],
  "Must","Corporate customers cannot prepay and reject PDF invoices; losing them is a documented risk.",
  [Q("CORP","Q3","One EHF invoice per month, with our cost-centre code on each line."),
   Q("EVC","Q3","Fjordtek and the other big firms want one invoice per month with all their lunches on it, sent as EHF."),
   Q("CSV","E004","our finance department only accepts EHF invoices with our PO number. Last invoice was a PDF by e-mail and was rejected."),
   Q("CSV","E016","Can we have separate logins but one monthly invoice?"),
   Q("ITL","Accounting","Fiken (has a REST API; supports EHF invoices)")],
  "DIRECT",PO+"Must is INFERRED from the documented risk of losing corporate customers. Different customers use different references (cost centre vs PO number)."),
 ("EP-04","As the owner, I want invoices for orders to be created in Fiken from the order data so that the coordinator no longer creates invoices manually.",
  ["Given an order is completed and invoiced, when the invoice is created, then it appears in Fiken with the order lines.",
   "Given an invoice created through the system, when it is viewed in Fiken, then it can be traced to its order number."],
  "Must","Kick-off states the system must integrate with Fiken; the owner will not change accounting system.",
  [Q("MIN","bullet 2","Must integrate with Fiken (accounting) and Vipps (payments)."),
   Q("OWN","Q3","We use Fiken for accounting today and I do not want to change that."),
   Q("ASIS","step 11","After the event the coordinator creates an invoice in Fiken manually"),
   Q("ASIS","step 12","Once a month the owner reconciles Fiken with the Excel sheet.")],
  "DIRECT","Fiken API coverage (invoice creation, payment status) must be confirmed (ISS-05). The traceability AC is an ASSUMPTION."),
 ("EP-05","As a corporate customer, I want a company account where several colleagues can log in and order so that we do not depend on one person sending e-mails.",
  ["Given a company account with several users, when any user places an order, then the order is linked to the company account.",
   "Given a user of the company account logs in, when they open the order list, then they see the company's orders, including standing orders."],
  "Should","Valuable for the largest customer segment; e-mail/coordinator entry still works at go-live.",
  [Q("CORP","Q2","A login for our company where I can see our standing orders"),
   Q("CORP","Q2","Several of my colleagues should be able to order"),
   Q("CSV","E016","Three of us order lunches for different departments. Can we have separate logins but one monthly invoice?")],
  "DIRECT","User administration (who adds users) and visibility between departments are not documented."),
 ("EP-05","As a corporate customer's office manager, I want orders above 15 000 NOK to require my approval so that large spending is controlled.",
  ["Given a user who is not the approver places an order above 15 000 NOK, when it is submitted, then it stays pending until the approver approves it.",
   "Given an order of 15 000 NOK or less, when it is submitted, then no approval is required."],
  "Could","One customer's internal control rule; not linked to go-live needs.",
  [Q("CORP","Q2","only I should approve anything above 15 000 NOK.")],
  "DIRECT",PO+"confirm whether the threshold is per customer and configurable (ISS-08). Whether 'above' includes VAT is unknown."),
 ("EP-05","As a corporate customer, I want to set up standing (recurring) orders so that I do not have to e-mail the same order every week.",
  ["Given a weekly standing order, when a new week starts, then an order for that week is created with the standing quantities.",
   "Given a company user logs in, when they open standing orders, then all active standing orders are listed."],
  "Should","Repeatedly requested by corporate customers; can be handled manually at go-live.",
  [Q("CORP","Q2","see our standing orders"),
   Q("CORP","Q1","I send an e-mail every Thursday with next week's lunch meetings."),
   Q("CSV","E007","Could we set up a standing order: lunch for 12 every Tuesday, vegetarian option for 3?")],
  "DIRECT","How far ahead standing orders are generated and how exceptions (holidays) are handled is not documented."),
 ("EP-05","As a returning customer, I want to copy a previous order so that I can reorder without re-entering the details.",
  ["Given a previous order, when the customer chooses to copy it, then a new draft order with the same items and guest groups is created for a new date."],
  "Should","Requested by two independent sources; not a go-live condition.",
  [Q("CORP","Q2","copy last week's order"),
   Q("CSV","E028","Could we just get exactly the same as our order in September? I don't have the details anymore.")],
  "DIRECT","Behaviour when a copied dish is no longer on the seasonal menu is not documented."),
 ("EP-05","As a corporate customer, I want our accountant to have read-only access to our invoices and order history so that they can check them without asking us.",
  ["Given a user with read-only access, when they log in, then they can view invoices and order history but cannot create or change orders."],
  "Could","Single request; low urgency.",
  [Q("CSV","E029","Can our accountant get read-only access to our invoices and order history?")],
  "DIRECT","Depends on US-25 (company accounts)."),
 ("EP-06","As a customer, I want to request access to the personal data Ekte Mat holds about me so that I can see what is stored.",
  ["Given a customer submits an access request, when the coordinator processes it, then the customer's personal data held in the system can be exported for the customer."],
  "Must","Documented obligation in the GDPR advice; applies from go-live because personal data is stored from day one.",
  [Q("GDPR","bullet 3","Customers must be able to request access and deletion.")],
  "DIRECT","Response deadline and identity verification are not documented (ISS-16). The export mechanism in the AC is an ASSUMPTION."),
 ("EP-06","As a customer, I want my personal data to be deleted when I ask, or when my company stops being a customer, so that my privacy is respected.",
  ["Given a deletion request, when it is processed, then the customer's personal data is deleted except data that must be kept for bookkeeping, and the customer receives a confirmation.",
   "Given a corporate customer relationship ends, when the deletion process runs, then the names of that company's employees are deleted."],
  "Must","Documented obligation in the GDPR advice and a corporate customer requirement.",
  [Q("GDPR","bullet 3","Customers must be able to request access and deletion."),
   Q("GDPR","bullet 2","Keep order data needed for bookkeeping 5 years (Norwegian Bookkeeping Act)"),
   Q("CSV","E010","Please delete all personal data you have about me and confirm by e-mail."),
   Q("CORP","Q4","they want to know that you delete our employees' names when we stop being a customer.")],
  "DIRECT","Deletion deadlines are missing (NFR-08, ISS-07). E010 is a real pending request in the current process; it is outside the system scope but should be handled now."),
]

# ---------------------------------------------------------------------- NFRs
NF = [
 ("Performance efficiency (time behaviour)","Customer-facing pages shall load fast enough not to deter customers. CANDIDATE: no target agreed.",
  "Page load time","TBC (not agreed)","s","To be defined once a target and test conditions (device, network, page set) are agreed.",
  [Q("CHF","Q4","Please do not make it slow. The old booking plugin on the website took ten seconds to load a page.")],
  "UNRESOLVED","10 s is the unacceptable baseline, not a target. Question in ISS-10."),
 ("Availability","Customers shall be able to view menus and send booking requests at any hour. CANDIDATE: availability target missing.",
  "Availability of menu and booking functions","TBC (not agreed)","% per month","To be defined: measurement period, maintenance windows, monitoring method.",
  [Q("OWN","Q2","send a booking request themselves, at any hour.")],
  "UNRESOLVED","'At any hour' gives no numeric target. Question in ISS-11."),
 ("Reliability (fault tolerance)","The kitchen view shall still show today's orders during an internet outage.",
  "Today's orders viewable on the kitchen device with the network disconnected","All of today's orders as of the last successful update","orders",
  "Test: load the kitchen view, disconnect the network, verify that all of today's orders remain visible.",
  [Q("CHF","Q3","It happened twice last winter. The kitchen must still see today's orders. Printing works too, but the screen should not just go blank.")],
  "DIRECT","Acceptable data age and whether changes are possible offline are not documented."),
 ("Security / privacy (data residency)","Personal data shall be stored and processed only in the EU/EEA.",
  "Share of personal data stored or processed in the EU/EEA","100","%","Review of hosting locations and sub-processors listed in each provider's data processing agreement.",
  [Q("ITL","Constraints","data must stay in EU/EEA (GDPR)"),
   Q("CORP","Q4","They want it in the EU/EEA")],
  "DIRECT","100 % follows from 'must stay'. Not a claim of legal compliance."),
 ("Privacy (data minimisation)","Allergy and dietary information shall be stored per guest group, not per named guest, unless strictly needed.",
  "Data fields linking allergy or diet data to a named guest","0 (unless a documented exception exists)","fields","Data-model review before release.",
  [Q("GDPR","bullet 1","Allergy and dietary information about named persons can be health data (special category, GDPR art. 9). Store per guest group, not per named guest, unless strictly needed.")],
  "DIRECT","Expressing this as a field count is an interpretation. 'Strictly needed' exceptions are undefined."),
 ("Privacy (retention)","Order data needed for bookkeeping shall be kept for 5 years.",
  "Retention period for bookkeeping order data","5","years","Review of retention configuration; test that bookkeeping data survives a deletion request.",
  [Q("GDPR","bullet 2","Keep order data needed for bookkeeping 5 years (Norwegian Bookkeeping Act)")],
  "DIRECT","Based on a phone note from the consultant; not legally verified here. Start of the period and the exact fields are undefined."),
 ("Privacy (retention)","Personal data not needed for bookkeeping shall have a shorter retention period. CANDIDATE: period missing.",
  "Retention period for other personal data","TBC (shorter than 5 years)","months","To be defined.",
  [Q("GDPR","bullet 2","other personal data should have a shorter retention period.")],
  "UNRESOLVED","Retention period not stated. Question in ISS-07."),
 ("Privacy (deletion)","Employee names of a corporate customer shall be deleted when the customer relationship ends. CANDIDATE: deadline missing.",
  "Time from end of customer relationship to deletion","TBC (not agreed)","days","To be defined.",
  [Q("CORP","Q4","they want to know that you delete our employees' names when we stop being a customer.")],
  "UNRESOLVED","Deadline and the definition of 'stop being a customer' are missing. Question in ISS-07."),
 ("Security / compliance (contracts)","A data processing agreement shall exist with every cloud or SaaS provider that processes personal data.",
  "Share of such providers with a signed data processing agreement","100","%","Contract register check before go-live.",
  [Q("GDPR","bullet 4","A data processing agreement is needed with any cloud/SaaS provider.")],
  "DIRECT","Applies to Vipps, Fiken, hosting and e-mail providers as relevant (INFERRED)."),
 ("Usability (localisation)","The user interface shall be available in Norwegian and English.",
  "Supported UI languages","2 (Norwegian, English)","languages","Review of all customer-facing screens in both languages.",
  [Q("ITL","Constraints","Norwegian and English UI."),
   Q("CSV","E012","Is your menu available in English with allergens marked?")],
  "DIRECT","Scope (customer-facing only, or also kitchen and admin views) is unresolved (ISS-19)."),
 ("Cost (operational constraint)","Running costs of the solution shall stay below 3 000 NOK per month.",
  "Monthly running cost","< 3 000","NOK/month","Sum of monthly hosting, licence and SaaS invoices after go-live.",
  [Q("BUD","Running cost","Running cost target: < 3 000 NOK/month"),
   Q("ITL","Constraints","prefer low running cost (< 3 000 NOK/month)")],
  "DIRECT","Project constraint rather than a product quality. Whether transaction fees (e.g. Vipps) count is not documented."),
 ("Maintainability / operability","Day-to-day content (e.g. menus) shall be manageable by non-technical staff. CANDIDATE: no measurable target.",
  "TBC","TBC (not agreed)","TBC","To be defined, e.g. through a usability test with staff once tasks and criteria are agreed.",
  [Q("ITL","Constraints","nobody technical in-house"),
   Q("CHF","Q2","I want to edit the menu myself")],
  "UNRESOLVED","No metric or threshold is documented."),
 ("Compatibility","The kitchen view shall run on a wall-mounted tablet.",
  "Supported kitchen device","TBC (device and OS not specified)","device","Test on the selected device.",
  [Q("CHF","Q1","Tablet on the wall."),
   Q("ITL","Kitchen","one old iPad nobody uses")],
  "INFERRED","Whether the existing iPad should be reused is not stated."),
]

# -------------------------------------------------------------------- issues
DO_PO = "Ingrid Solheim (product owner, per 00_minutes bullet 5)"
IS = [
 ("Conflict","Deposit / prepayment","The owner wants a 30 % deposit on every booking; the event coordinator and the corporate customer say corporate customers cannot or will not prepay.",
  "Owner (01_interview_owner): 30 % deposit for every booking, 'no deposit, no booking'.",
  "Event coordinator and corporate customer: no prepayment for corporate customers; deposit acceptable for weddings.",
  [Q("OWN","Q2","I want the system to take a 30 % deposit when someone books - no deposit, no booking."),
   Q("EVC","Q3","If we force a deposit on them, we will lose them. For weddings a deposit makes sense."),
   Q("CORP","Q3","We cannot prepay, our finance department does not allow it.")],
  "Affects US-22 and US-23. Risk of losing corporate customers, or of continued no-shows.",
  "Which booking types require the 30 % deposit: (a) all, (b) weddings only, (c) all except corporate customers on monthly invoicing? What are the refund rules?",DO_PO),
 ("Open Question","Build vs buy","Kick-off left open whether to build a custom system or configure a catering SaaS.",
  "Kick-off minutes: open question.","-",
  [Q("MIN","bullet 8","Open question: build our own system or configure an off-the-shelf catering SaaS?")],
  "Affects architecture, cost (NFR-11), data processing agreements (NFR-09) and all stories.",
  "Will the team build its own system or configure a catering SaaS? By which sprint must this be decided?","Not evidenced"),
 ("Open Question","Order change deadline","No company-wide deadline for customer changes exists. One customer wants changes until 10:00 the day before; the vision says 'change orders any time'.",
  "Corporate customer: changes until 10:00 the day before.","Vision draft v0.3: 'book and change orders any time'.",
  [Q("CORP","Q2","change the number of people until 10:00 the day before"),
   Q("VIS","para 1","book and change orders any time")],
  "Affects US-11 and kitchen planning.",
  "What is the latest time a customer may change guest numbers or allergies, and does it differ for lunches and events?",DO_PO),
 ("Missing Evidence","Cancellation terms","A customer asks for cancellation terms; none are documented.",
  "Customer e-mail E013: asks for terms.","-",
  [Q("CSV","E013","What are your cancellation terms?")],
  "US-13 cannot be fully specified.","What are the cancellation deadlines and fees for each booking type?",DO_PO),
 ("Dependency","Fiken and EHF","Invoice creation, EHF sending and payment-status updates depend on Fiken's API. A customer received a wrong payment reminder.",
  "IT landscape: Fiken has a REST API and supports EHF.","Customer e-mail E030: wrong payment reminder.",
  [Q("ITL","Accounting","Fiken (has a REST API; supports EHF invoices)"),
   Q("CSV","E030","We received a payment reminder for an invoice we paid on 2 Sept.")],
  "Affects US-23 and US-24.",
  "Can the team get test access to Fiken's API, and should payment status be read back from Fiken into the system?",DO_PO),
 ("Open Question","Capacity rule","Capacity is described as 'about 400 covers a day, 250 if there is a wedding with a served dinner'. It is unclear whether this is a hard limit or a warning.",
  "Event coordinator: approximate limits.","-",
  [Q("EVC","Q4","Our kitchen can do about 400 covers a day, 250 if there is a wedding with a served dinner.")],
  "Affects US-03 and US-10.",
  "Are 400 and 250 covers per day hard limits that block bookings, or warnings to the coordinator? Who may override them?",DO_PO),
 ("Missing Evidence","Retention and deletion periods","Retention for non-bookkeeping personal data and deadlines for deletion are not stated.",
  "GDPR note: 'shorter retention period' (no value).","Corporate customer: deletion at end of relationship (no deadline).",
  [Q("GDPR","bullet 2","other personal data should have a shorter retention period."),
   Q("CORP","Q4","delete our employees' names when we stop being a customer.")],
  "NFR-07 and NFR-08 cannot be tested; affects US-31.",
  "What retention period applies to personal data not needed for bookkeeping, and within how many days must data be deleted after a request or end of relationship?","Not evidenced"),
 ("Open Question","Approval threshold","One customer wants approval above 15 000 NOK. It is unclear whether thresholds are per customer and configurable.",
  "Corporate customer: only the office manager approves orders above 15 000 NOK.","-",
  [Q("CORP","Q2","only I should approve anything above 15 000 NOK.")],
  "Affects US-26.","Should approval thresholds be configurable per company account, and is the amount incl. or excl. VAT?",DO_PO),
 ("Dependency","Vipps integration","Vipps payments depend on an agreement and API access that are not documented.",
  "Kick-off: must integrate with Vipps.","-",
  [Q("MIN","bullet 2","Must integrate with Fiken (accounting) and Vipps (payments).")],
  "Affects US-21 and US-22.","Does Ekte Mat have a Vipps merchant agreement that allows API integration, and who will request test access?",DO_PO),
 ("Missing Evidence","Performance target","Speed is a stated need, but no target is given.",
  "Head chef: 'do not make it slow'.","-",
  [Q("CHF","Q4","Please do not make it slow.")],
  "NFR-01 cannot be tested.","What is the maximum acceptable page load time, on which device and network?",DO_PO),
 ("Missing Evidence","Availability target","'At any hour' has no numeric availability target.",
  "Owner: booking at any hour.","-",
  [Q("OWN","Q2","at any hour.")],
  "NFR-02 cannot be tested; affects hosting choice and cost.","What monthly availability (%) is required for menus and booking, and are planned maintenance windows acceptable?",DO_PO),
 ("Open Question","Phone ordering","Phone orders are part of today's process and a customer asks if phone ordering continues.",
  "As-is process: phone calls noted on paper.","Customer e-mail E024: asks to keep ordering by phone.",
  [Q("ASIS","step 1","Customer sends an e-mail or calls (phone calls are noted on paper)."),
   Q("CSV","E024","Kan jeg fortsatt bestille på telefon?","Can I still order by phone?")],
  "Affects US-09 and the scope of the coordinator's role.","Will phone and e-mail orders still be accepted after go-live, with the coordinator entering them in the system?",DO_PO),
 ("Missing Evidence","Driver needs","The driver is part of the process but was not interviewed.",
  "As-is process / IT landscape: driver uses printouts and own phone.","-",
  [Q("ASIS","step 10","Driver gets the address from the printout."),
   Q("ITL","Driver","Driver: own phone, Google Maps")],
  "Delivery stories may be incomplete (US-14).","Should the team interview the driver before Sprint 2, and does the driver need access to the system?",DO_PO),
 ("Open Question","Kitchen PIN vs health data","A PIN login is requested for the kitchen tablet, which will display allergy data that may be health data.",
  "Head chef: PIN is fine.","GDPR note: allergy data about named persons can be health data.",
  [Q("CHF","Q1","a PIN is fine."),
   Q("GDPR","bullet 1","Allergy and dietary information about named persons can be health data")],
  "Affects US-17 and NFR-05. Not a conflict if allergy data is stored per guest group.","Is a shared or personal PIN acceptable for the kitchen tablet, given that it shows allergy information per guest group?",DO_PO),
 ("Missing Evidence","Recipe quantities","A prep list in kilos requires ingredient quantities per portion, which are not in the evidence.",
  "Head chef: prep list in kilos.","-",
  [Q("CHF","Q2","total kilos of each ingredient for the next three days.")],
  "US-20 cannot be implemented without recipe data.","Who will provide ingredient quantities per portion for each dish, and in what format?",DO_PO),
 ("Missing Evidence","Data-subject request handling","Deadline and identity checks for access and deletion requests are not documented. A deletion request (E010) is already pending.",
  "GDPR note: customers must be able to request access and deletion.","-",
  [Q("GDPR","bullet 3","Customers must be able to request access and deletion."),
   Q("CSV","E010","Please delete all personal data you have about me and confirm by e-mail.")],
  "Affects US-30 and US-31.","Who handles access and deletion requests, how is identity verified, and what response deadline applies? Has request E010 been handled?","Not evidenced"),
 ("Missing Evidence","Pricing rules","Customers ask about volume discounts and a children's menu; no pricing rules are documented.",
  "Customer e-mail E020: discount above 150 guests.","Customer e-mail E027: children's menu.",
  [Q("CSV","E020","Do you have a discount above 150 guests?"),
   Q("CSV","E027","Do you have a children's menu?")],
  "Affects price display (US-01) and quotes.","Are there volume discounts or special menus (e.g. children) that the system must calculate or display?",DO_PO),
 ("Missing Evidence","Accessibility","No accessibility requirement or standard is documented. An elderly customer prefers phone ordering.",
  "Customer e-mail E024.","-",
  [Q("CSV","E024","Jeg er 82 år og bruker ikke internett så mye.","I am 82 and do not use the internet much.")],
  "No accessibility NFR could be written.","Must the customer-facing website meet a specific accessibility standard (e.g. WCAG level)?",DO_PO),
 ("Open Question","Language scope","Norwegian and English UI is required, but the scope is not defined.",
  "IT landscape: Norwegian and English UI.","-",
  [Q("ITL","Constraints","Norwegian and English UI.")],
  "Affects NFR-10 and effort.","Must kitchen and admin screens also be available in both languages, or only customer-facing pages and menus?",DO_PO),
 ("Open Question","Go-live meaning","Sources describe 2027-03-01 as a 'soft launch target' and as a 'hard' go-live. Whether a soft launch allows a reduced scope is unclear.",
  "Kick-off minutes: soft launch target.","Budget note: go-live hard.",
  [Q("MIN","bullet 4","Soft launch target: 1 March 2027."),
   Q("BUD","Go-live","Go-live: 1 March 2027 (hard)")],
  "Determines which stories are Must.","Is 2027-03-01 a hard date for the full Must scope, or a soft launch with a limited customer group?",DO_PO),
 ("Open Question","Existing tools","It is unclear whether the new system replaces the Wix website and the Google Calendar capacity check, or integrates with them.",
  "IT landscape: Wix website, Google Workspace.","-",
  [Q("ITL","Website","Website: Wix, 6 pages, contact form -> sends e-mail to post@ektemat.example"),
   Q("MIN","bullet 2","Google Workspace used for e-mail/calendar")],
  "Affects architecture and US-02, US-10.","Will the new system replace the Wix website and the Google Calendar capacity check, or must it integrate with them?",DO_PO),
 ("Open Question","Scope vs budget","The Must scope (11 stories) has not been estimated against the 450 000 NOK budget and the go-live date.",
  "Budget note.","-",
  [Q("BUD","Budget","Budget first version: 450 000 NOK incl. 12 months hosting")],
  "Risk that the Must scope does not fit.","After estimation, which Must stories may be moved to Should if the scope exceeds budget or time?",DO_PO),
 ("Open Question","Small epic","EP-06 (Privacy and data-subject rights) has only 2 stories, below the 3-story guideline. No stories were invented to fill it.",
  "Backlog rule: 3-8 stories per epic.","-",
  [Q("GDPR","bullet 3","Customers must be able to request access and deletion.")],
  "Backlog structure only.","Keep EP-06 as a small compliance epic, or merge its stories into another epic?",DO_PO),
 ("Missing Evidence","GDPR advice verification","GDPR obligations come from a phone note, not a written legal assessment.",
  "GDPR note header: phone call.","-",
  [Q("GDPR","header","Note from data-protection consultant (phone call, 20 Aug 2026)")],
  "Risk of building on unverified legal assumptions (NFR-04 to NFR-09).","Can the consultant confirm the advice in writing, including the 5-year bookkeeping retention?","Not evidenced"),
]

# ---------------------------------------------------------------- validation
errors = []
def check_quotes(where, quotes):
    for k, ref, t, tr in quotes:
        hay = EMAIL.get(ref, "") if k == "CSV" else TEXT[k]
        if k == "CSV" and ref not in EMAIL: errors.append(f"{where}: unknown email_id {ref}")
        if norm(t) not in hay: errors.append(f"{where}: quote not found in {k} [{ref}]: {t}")
for i, s in enumerate(STAKE, 1): check_quotes(f"SH-{i:02d}", s[5])
for i, s in enumerate(ST, 1): check_quotes(f"US-{i:02d}", s[5])
for i, s in enumerate(NF, 1): check_quotes(f"NFR-{i:02d}", s[6])
for i, s in enumerate(IS, 1): check_quotes(f"ISS-{i:02d}", s[5])

for i, s in enumerate(ST, 1):
    if not (1 <= len(s[2]) <= 3): errors.append(f"US-{i:02d}: {len(s[2])} ACs")
    if s[3] not in ("Must", "Should", "Could", "Won't"): errors.append(f"US-{i:02d}: bad MoSCoW")
    if not s[4]: errors.append(f"US-{i:02d}: no rationale")
    if not re.match(r"As (a|an|the) ", s[1]) or " I want " not in s[1] or " so that " not in s[1]:
        errors.append(f"US-{i:02d}: story format")
    for ac in s[2]:
        if not (ac.startswith("Given ") and " when " in ac and " then " in ac): errors.append(f"US-{i:02d}: AC format")

epic_sizes = collections.Counter(s[0] for s in ST)
order = [s[0] for s in ST]
if order != sorted(order): errors.append("stories not grouped by epic order")
small = [e for e, n in epic_sizes.items() if n < 3]
big = [e for e, n in epic_sizes.items() if n > 8]
if big: errors.append(f"epics > 8: {big}")
if small != ["EP-06"]: errors.append(f"unexpected small epics {small}")
if "Small epic" not in [x[1] for x in IS]: errors.append("small epic not flagged")

# Cross-references to IDs that must exist
refs = re.findall(r"\b(US|NFR|ISS)-(\d\d)\b", " ".join(str(c) for r in ST + NF + IS for c in r))
lim = {"US": len(ST), "NFR": len(NF), "ISS": len(IS)}
for p, n in refs:
    if int(n) > lim[p]: errors.append(f"dangling reference {p}-{n}")

per_file = collections.Counter()
for s in ST:
    for f in set(sources(s[5]).split("\n")): per_file[f] += 1

n_must = sum(1 for s in ST if s[3] == "Must")
if f"The Must scope ({n_must} stories)" not in " ".join(x[2] for x in IS): errors.append("Must count in ISS text out of date")
if errors:
    print("VALIDATION FAILED"); print("\n".join(errors)); sys.exit(1)

# -------------------------------------------------------------------- writer
wb = Workbook()
F = Font(name="Arial", size=10); FH = Font(name="Arial", size=10, bold=True, color="FFFFFF")
FILL = PatternFill("solid", fgColor="1F4E78")
WRAP = Alignment(wrap_text=True, vertical="top")

def write_table(ws, start_row, headers, rows, name, widths):
    for c, h in enumerate(headers, 1):
        cell = ws.cell(start_row, c, h); cell.font = FH; cell.fill = FILL; cell.alignment = WRAP
    for r, row in enumerate(rows, start_row + 1):
        for c, v in enumerate(row, 1):
            cell = ws.cell(r, c, v); cell.font = F; cell.alignment = WRAP
    end = start_row + len(rows)
    ref = f"A{start_row}:{get_column_letter(len(headers))}{end}"
    t = Table(displayName=name, ref=ref)
    t.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
    ws.add_table(t)
    for c, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(c)].width = w
    return end

def dv(ws, options, col, r1, r2):
    d = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=False)
    d.add(f"{col}{r1}:{col}{r2}"); ws.add_data_validation(d)

EV = ["DIRECT", "INFERRED", "ASSUMPTION", "UNRESOLVED"]

ws = wb.active; ws.title = "Stakeholders"
H1 = ["Stakeholder ID","Name/Group","Role","Interest","Influence","Key Concerns","Source File(s)","Supporting Quote/Evidence","Evidence Status"]
rows = [[f"SH-{i:02d}", s[0], s[1], s[2], s[3], s[4], sources(s[5]), render(s[5]) + "\nNote: " + s[7], s[6]] for i, s in enumerate(STAKE, 1)]
e = write_table(ws, 1, H1, rows, "Stakeholders", [10,24,28,40,10,40,40,60,13])
dv(ws, ["High","Medium","Low"], "E", 2, e); dv(ws, EV, "I", 2, e); ws.freeze_panes = "B2"

ws = wb.create_sheet("User stories")
H2 = ["Story ID","Epic ID","Epic Name","User Story","Acceptance Criteria (2–3)","MoSCoW Priority","One-line Priority Rationale","Source File(s)","Supporting Quote","Evidence Status","Assumptions/Dependencies"]
rows = []
for i, s in enumerate(ST, 1):
    ac = "\n".join(f"AC{j}: {a}" for j, a in enumerate(s[2], 1))
    if len(s[2]) < 2: ac += "\n[Only one AC is supported by the evidence; further criteria need PO input.]"
    rows.append([f"US-{i:02d}", s[0], EPICS[s[0]], s[1], ac, s[3], s[4], sources(s[5]), render(s[5]), s[6], s[7]])
e = write_table(ws, 1, H2, rows, "UserStories", [9,8,22,45,60,10,38,40,60,13,45])
dv(ws, ["Must","Should","Could","Won't"], "F", 2, e); dv(ws, EV, "J", 2, e); ws.freeze_panes = "B2"
STORY_SRC_RANGE = f"'User stories'!$H$2:$H${e}"; STORY_ID_RANGE = f"'User stories'!$A$2:$A${e}"

ws = wb.create_sheet("Non-functional requirements")
H3 = ["Requirement ID","Quality Characteristic","Requirement Statement","Metric","Target Value","Unit","Measurement Conditions/Verification Method","Source File(s)","Supporting Quote","Evidence Status","Assumptions/Unresolved Details"]
rows = [[f"NFR-{i:02d}", s[0], s[1], s[2], s[3], s[4], s[5], sources(s[6]), render(s[6]), s[7], s[8]] for i, s in enumerate(NF, 1)]
e = write_table(ws, 1, H3, rows, "NFRs", [11,24,45,30,20,11,40,40,60,13,45])
dv(ws, EV, "J", 2, e); ws.freeze_panes = "B2"

ws = wb.create_sheet("Conflicts and open questions")
H4 = ["Issue ID","Issue Type","Topic","Description","Stakeholder/Source A and Position","Stakeholder/Source B and Position","Source File(s)","Supporting Quote/Evidence","Potential Impact","Specific Product-Owner Question","Decision Owner (if evidenced)","Status"]
rows = [[f"ISS-{i:02d}", s[0], s[1], s[2], s[3], s[4], sources(s[5]), render(s[5]), s[6], s[7], s[8], "Open"] for i, s in enumerate(IS, 1)]
e = write_table(ws, 1, H4, rows, "Issues", [9,15,20,45,35,35,40,60,35,50,26,10])
dv(ws, ["Conflict","Open Question","Missing Evidence","Dependency"], "B", 2, e)
dv(ws, ["Open","In review","Resolved"], "L", 2, e); ws.freeze_panes = "B2"

# Audit trail
r0 = e + 3
ws.cell(r0, 1, "AUDIT TRAIL (separate from the issue table above)").font = Font(name="Arial", size=12, bold=True)
ws.cell(r0 + 1, 1, "Inspected 2026-10-22. Story counts are live formulas over the 'User stories' Source File(s) column.").font = F
HA = ["#","File","Relative Path (from EkteMat_repo)","File Type","Read Successfully","Contributed Stories","Distinct Stories Supported","Other Contribution","Notes / Exclusion Reason"]
AUD = [
 (SRC["OWN"],"txt","Yes","Stakeholders, NFR-02, issues",""),
 (SRC["EVC"],"txt","Yes","Stakeholders, issues",""),
 (SRC["CHF"],"txt","Yes","NFR-01, 03, 12, 13; issues",""),
 (SRC["CORP"],"txt","Yes","NFR-04, 08; issues",""),
 (SRC["GDPR"],"txt","Yes","NFR-05 to 09; issues","Advice from a phone note; not legally verified."),
 (SRC["VIS"],"docx","Yes","Issues","Text read from paragraphs; no tables or images found."),
 (SRC["MIN"],"txt","Yes","Stakeholders, issues","Expected input; located in 00_Project_management, not in sources/."),
 (SRC["BUD"],"txt","Yes","NFR-11, issues","Expected input; located in 00_Project_management. No functional stories: content is constraints only."),
 (SRC["ASIS"],"md","Yes","Stakeholders, issues","Expected input; located in 02_Analysis_and_models."),
 (SRC["ITL"],"txt","Yes","NFR-04, 10-13; issues","Expected input; located in 03_Architecture."),
 (SRC["CSV"],"csv","Yes","Stakeholders, NFR-10, issues","Expected input; located in 04_Data. 30 records; columns email_id, received, from_name, from_address, subject, body, language. Cited by email_id; names and addresses not reproduced. Records used: see note below."),
 (SRC["MENU"],"txt","Yes","-","Reference material (allergen list). Not in the expected-input list; included because it supports US-06."),
 ("04_Data/04_dataset_orders-2024-2025_2026-09-27.xlsx","xlsx","Partially (headers only)","-","Historical order data; not a requirements source in the brief. Not analysed for stories."),
 ("02_Analysis_and_models/02_benchmark_competitor-links_2026-10-20.txt","txt","Yes","-","Excluded: competitor links for benchmarking only; external sites not opened (no research authorised)."),
 ("00_Project_management/00_todo_sprint0-setup_2026-10-20.txt","txt","Yes","-","Excluded: team to-do, no requirements content."),
 ("99_Not_relevant/99_note_logo-ideas_2026-10-20.txt","txt","Yes","-","Excluded: marked not relevant to system development."),
 ("README.md","md","Yes","-","Excluded: repository index written by the team."),
]
rows = []
for i, (rel, typ, ok, other, note) in enumerate(AUD, 1):
    rr = r0 + 2 + i
    name = Path(rel).name
    rows.append([i, name, rel, typ, ok, f'=IF(G{rr}>0,"Yes","No")', f'=COUNTIF({STORY_SRC_RANGE},"*"&B{rr}&"*")', other, note])
ea = write_table(ws, r0 + 2, HA, rows, "AuditTrail", [9,15,20,45,35,35,40,60,35])
for c, w in enumerate([5,40,55,9,14,12,14,30,60], 1):
    ws.column_dimensions[get_column_letter(c)].width = max(ws.column_dimensions[get_column_letter(c)].width or 0, w)
t1 = ea + 2
ws.cell(t1, 2, "Total unique user stories").font = Font(name="Arial", size=10, bold=True)
ws.cell(t1, 7, f"=COUNTA({STORY_ID_RANGE})").font = F
ws.cell(t1 + 1, 2, "Sum of per-file story counts").font = Font(name="Arial", size=10, bold=True)
ws.cell(t1 + 1, 7, f"=SUM(G{r0+3}:G{ea})").font = F
ws.cell(t1 + 2, 2, "Note").font = Font(name="Arial", size=10, bold=True)
csv_used = sorted({ref for s in ST for k, ref, *_ in s[5] if k == "CSV"})
other_q = [q for s in STAKE for q in s[5]] + [q for s in NF for q in s[6]] + [q for s in IS for q in s[5]]
csv_other = sorted({q[1] for q in other_q if q[0] == "CSV"} - set(csv_used))
csv_none = sorted(set(EMAIL) - set(csv_used) - set(csv_other))
note = ("The sum exceeds the unique total because one story can be supported by several files. "
        f"CSV records supporting stories: {', '.join(csv_used)}. "
        f"Records used only for stakeholders/NFRs/issues: {', '.join(csv_other) or 'none'}. "
        f"Records with no requirement content (not used): {', '.join(csv_none) or 'none'}. "
        "No story is classified Won't: no source explicitly excludes or defers a need. "
        "Originals in EkteMat_project_inbox were not re-read; the repo copies are SHA-256-identical (verified 2026-10-22). "
        "The inbox file 'Kopi av interview Sunniva (2).txt' is an exact duplicate of the event-coordinator interview and was not counted.")
c = ws.cell(t1 + 2, 3, note); c.font = F; c.alignment = WRAP
ws.merge_cells(start_row=t1 + 2, start_column=3, end_row=t1 + 2, end_column=9)
ws.row_dimensions[t1 + 2].height = 105

OUT.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUT)
print("OK", OUT)
print("stakeholders", len(STAKE), "stories", len(ST), "epics", len(epic_sizes), dict(sorted(epic_sizes.items())),
      "nfrs", len(NF), "issues", len(IS))
print("moscow", collections.Counter(s[3] for s in ST))
print("per-file", dict(per_file), "sum", sum(per_file.values()))
print("csv used", csv_used, "other", csv_other, "none", csv_none)
